"""Excel processing service - extracted from main_window.py"""
import logging
import openpyxl
import glob
import os
from pathlib import Path
from typing import Optional, List
from dataclasses import dataclass
import pandas as pd
from modules.excel_handler import ExcelBebanProcessor

logger = logging.getLogger(__name__)

@dataclass
class ExcelResult:
    """Result of Excel operation"""
    success: bool
    message: str
    rows_processed: int = 0
    next_id: Optional[int] = None
    error: Optional[str] = None


def copy_excel_data(
    rencana_path: str,
    pengujian_path: str,
    sheet_source: str,
    sheet_odoo: str,
    sheet_pengujian: str
) -> ExcelResult:
    """
    Copy data from Rencana to Pengujian Excel, calculate BEBAN, update idpengujian.
    
    Args:
        rencana_path: Path to Rencana Excel file
        pengujian_path: Path to Pengujian Excel file
        sheet_source: Source sheet name
        sheet_odoo: ODOO sheet name
        sheet_pengujian: Pengujian sheet name
    
    Returns:
        ExcelResult with success status and details
    """
    try:
        # Column names to copy
        kolom_yang_dicopy = [
            'Docket', 'Nomor Kontrak', 'Nomor SPP', 'Proyek', 'Kontraktor', 
            'Mutu', 'No. Urut', 'Tanggal Rencana', 'Rencana Umur Test', 
            'Bentuk Benda Uji', 'Target (%)', 'Tanggal Realisasi Test', 
            'Realisasi Umur Test', 'Hasil Test', 'Rusak', 'Kode Benda Uji', 
            'Jenis Retakan', 'Kn', 'Mpa', 'Kg/cm2', 'Berat (kg)', 
            'Persentase Kekuatan (%)'
        ]
        
        # Load source workbook
        logger.info("Loading source Excel file...")
        wb1 = openpyxl.load_workbook(rencana_path)
        sheet1 = wb1[sheet_source]
        
        # Get header and column indices
        header = [cell.value for cell in sheet1[1]]
        index_kolom = [header.index(k) + 1 for k in kolom_yang_dicopy if k in header]
        
        logger.info(f"Found {len(index_kolom)} columns to copy")
        
        # Extract data from specified columns
        data_kolom = []
        for col in index_kolom:
            col_data = []
            for row in sheet1.iter_rows(min_row=2, min_col=col, max_col=col):
                val = row[0].value
                col_data.append(val)
            data_kolom.append(col_data)
        
        total_rows = len(data_kolom[0]) if data_kolom else 0
        logger.info(f"Total rows to process: {total_rows}")
        
        # Load target workbook
        logger.info("Loading target Excel file...")
        wb2 = openpyxl.load_workbook(pengujian_path)
        sheet2 = wb2[sheet_odoo]
        
        # Write data to target sheet
        baris_mulai = 2
        kolom_mulai = 1
        
        logger.info("Copying data to target Excel...")
        for i, kolom_data in enumerate(data_kolom):
            for j, value in enumerate(kolom_data):
                sheet2.cell(row=baris_mulai + j, column=kolom_mulai + i, value=value)
        
        # Save target workbook before BEBAN processing
        wb2.save(pengujian_path)
        wb2.close()
        
        # Calculate BEBAN
        logger.info("Calculating BEBAN values...")
        beban_processor = ExcelBebanProcessor(pengujian_path, sheet_odoo)
        if beban_processor.load_excel():
            beban_processor.process_all_rows(start_row=2)
            
            if beban_processor.save_excel():
                logger.info("BEBAN values calculated and saved successfully")
            else:
                logger.error("Failed to save BEBAN calculations")
            
            beban_processor.close()
        else:
            logger.error("Failed to load Excel file for BEBAN processing")
        
        # Update idpengujian
        logger.info("Updating idpengujian...")
        next_id = update_idpengujian(pengujian_path, sheet_pengujian)
        
        logger.info("All processes completed successfully!")
        
        return ExcelResult(
            success=True,
            message=f"Copied {total_rows} rows, calculated BEBAN, updated idpengujian to {next_id}",
            rows_processed=total_rows,
            next_id=next_id
        )
        
    except FileNotFoundError as e:
        logger.error(f"File not found: {str(e)}")
        return ExcelResult(
            success=False,
            message=f"File not found: {str(e)}",
            error=str(e)
        )
    except KeyError as e:
        logger.error(f"Sheet not found: {str(e)}")
        return ExcelResult(
            success=False,
            message=f"Sheet not found: {str(e)}",
            error=str(e)
        )
    except Exception as e:
        logger.error(f"Error during copy_excel_data: {str(e)}")
        return ExcelResult(
            success=False,
            message=f"Failed: {str(e)}",
            error=str(e)
        )


def update_idpengujian(pengujian_path: str, sheet_name: str, output_dir: str = "Output") -> int:
    """
    Update idpengujian in pengujian sheet based on latest CSV.
    
    Args:
        pengujian_path: Path to Pengujian Excel file
        sheet_name: Pengujian sheet name
        output_dir: Output directory for CSVs
    
    Returns:
        Next idpengujian value
    """
    # Find latest CSV
    csv_files = glob.glob(os.path.join(output_dir, "*.csv"))
    if csv_files:
        latest_csv = max(csv_files, key=os.path.getctime)
        df_latest = pd.read_csv(latest_csv, sep=';', encoding='utf-8-sig')
        if 'idpengujian' in df_latest.columns and not df_latest['idpengujian'].empty:
            last_id = pd.to_numeric(df_latest['idpengujian'], errors='coerce').dropna().astype(int).max()
            next_id = last_id + 1
        else:
            next_id = 1
    else:
        next_id = 1
    
    # Update idpengujian value in pengujian sheet
    wb_pengujian = openpyxl.load_workbook(pengujian_path)
    ws_pengujian = wb_pengujian[sheet_name]
    
    # Find idpengujian column
    for col in range(1, ws_pengujian.max_column + 1):
        if ws_pengujian.cell(row=1, column=col).value == 'idpengujian':
            ws_pengujian.cell(row=2, column=col).value = next_id
            break
    
    wb_pengujian.save(pengujian_path)
    wb_pengujian.close()
    logger.info(f"Updated idpengujian to: {next_id}")
    
    return next_id
