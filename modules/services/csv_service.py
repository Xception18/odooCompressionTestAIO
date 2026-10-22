"""CSV generation service - extracted from main_window.py"""
import logging
import openpyxl
import pandas as pd
import xlwings as xw
import os
from dataclasses import dataclass
from typing import List

logger = logging.getLogger(__name__)

@dataclass
class CsvResult:
    """Result of CSV generation"""
    success: bool
    message: str
    files_created: List[str] = None
    error: str = None
    
    def __post_init__(self):
        if self.files_created is None:
            self.files_created = []


def generate_csv_files(
    pengujian_path: str,
    sheet_pengujian: str,
    sheet_odoo: str,
    sheet_bjdt: str,
    output_dir: str
) -> CsvResult:
    """
    Generate CSV files grouped by date and clear sheets.
    
    Args:
        pengujian_path: Path to Pengujian Excel file
        sheet_pengujian: Pengujian sheet name
        sheet_odoo: ODOO sheet name to clear
        sheet_bjdt: BJDT ID sheet name to clear
        output_dir: Output directory
    
    Returns:
        CsvResult with success status and file list
    """
    try:
        # Calculate formulas
        logger.info("Calculating formulas...")
        with xw.App(visible=False) as app:
            wb = app.books.open(pengujian_path)
            wb.app.calculate()
            wb.save()
            wb.close()
        
        # Read Excel file
        logger.info("Reading Excel data...")
        df = pd.read_excel(pengujian_path, sheet_name=sheet_pengujian, engine='openpyxl')
        
        # Convert date column to datetime
        df['tglrencanauji'] = pd.to_datetime(df['tglrencanauji'], errors='coerce')
        
        # Create output directory
        os.makedirs(output_dir, exist_ok=True)
        
        # Group by date and save as CSV
        files_created = []
        for tgl, group in df.groupby(df['tglrencanauji'].dt.date):
            output_filename = os.path.join(output_dir, f"{tgl}.csv")
            group.to_csv(output_filename, index=False, sep=';', encoding='utf-8-sig')
            logger.info(f"Saved: {output_filename}")
            files_created.append(output_filename)
        
        # Clear ODOO sheet
        logger.info("Clearing ODOO sheet...")
        wb = openpyxl.load_workbook(pengujian_path)
        ws = wb[sheet_odoo]
        
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                cell.value = None
        
        wb.save(pengujian_path)
        
        # Clear bjdt_id sheet (columns 3 and beyond)
        logger.info("Clearing BJDT ID sheet...")
        ws = wb[sheet_bjdt]
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=3, max_col=ws.max_column):
            for cell in row:
                cell.value = None
        
        wb.save(pengujian_path)
        wb.close()
        
        logger.info(f"CSV files generated: {len(files_created)}")
        
        return CsvResult(
            success=True,
            message=f"Generated {len(files_created)} CSV files",
            files_created=files_created
        )
        
    except Exception as e:
        logger.error(f"Error during CSV generation: {str(e)}")
        return CsvResult(
            success=False,
            message=f"Failed: {str(e)}",
            error=str(e)
        )
