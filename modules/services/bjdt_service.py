"""BJDT ID fetching service - extracted from main_window.py"""
import logging
import openpyxl
import requests
import xlwings as xw
from dataclasses import dataclass
from typing import List, Optional

logger = logging.getLogger(__name__)

@dataclass
class BjdtResult:
    """Result of BJDT ID fetching"""
    success: bool
    message: str
    ids_found: int = 0
    ids_missing: int = 0
    errors: List[str] = None
    
    def __post_init__(self):
        if self.errors is None:
            self.errors = []


def get_bjdt_ids(
    pengujian_path: str,
    sheet_name: str,
    timeout: int = 10
) -> BjdtResult:
    """
    Get BJDT IDs from API and update Excel.
    
    Args:
        pengujian_path: Path to Pengujian Excel file
        sheet_name: Sheet name (bjdt_id)
        timeout: Request timeout in seconds
    
    Returns:
        BjdtResult with success status and statistics
    """
    try:
        # Calculate formulas first
        logger.info("Calculating formulas...")
        with xw.App(visible=False) as app:
            wb = app.books.open(pengujian_path)
            wb.app.calculate()
            wb.save()
            wb.close()
        
        # Load workbooks
        wb_data = openpyxl.load_workbook(pengujian_path, data_only=True)
        sheet_data = wb_data[sheet_name]
        
        wb_write = openpyxl.load_workbook(pengujian_path)
        sheet_write = wb_write[sheet_name]
        
        # Get header and column indices
        header = [cell.value for cell in sheet_data[1]]
        
        try:
            index_doc_no = header.index('Docket')
            index_no_urut = header.index('No URUT')
        except ValueError as e:
            logger.error(f"Required column not found: {e}")
            return BjdtResult(
                success=False,
                message=f"Required column not found: {e}",
                errors=[str(e)]
            )
        
        # Add ID column if not exists
        if 'ID' not in header:
            sheet_write.cell(row=1, column=len(header) + 1, value='ID')
            index_bjdt_id = len(header)
        else:
            index_bjdt_id = header.index('ID')
        
        error_logs = []
        ids_found = 0
        ids_missing = 0
        
        # Process each row
        for i, (row_data, row_write) in enumerate(zip(
            sheet_data.iter_rows(min_row=2), 
            sheet_write.iter_rows(min_row=2)
        ), start=2):
            
            doc_no = row_data[index_doc_no].value
            no_urut = row_data[index_no_urut].value
            
            if doc_no and no_urut:
                doc_no = str(doc_no).strip()
                no_urut = str(no_urut).strip()
                url = f"https://rmc.adhimix.web.id/benda_uji/?doc_no={doc_no}&no_urut={no_urut}"
                
                try:
                    logger.info(f"[{i}] Getting bjdt_id for Doc: {doc_no}, No Urut: {no_urut}...")
                    response = requests.get(url, timeout=timeout)
                    response.raise_for_status()
                    data = response.json()
                    bjdt_id = data.get("bjdt_id", "")
                    
                    if not bjdt_id:
                        logger.warning(f"bjdt_id NOT FOUND")
                        error_logs.append(f"bjdt_id not found for {doc_no} no_urut {no_urut}")
                        ids_missing += 1
                    else:
                        logger.info(f"Found bjdt_id: {bjdt_id}")
                        ids_found += 1
                    
                    row_write[index_bjdt_id].value = bjdt_id
                    
                except requests.Timeout:
                    logger.error(f"TIMEOUT on {url}")
                    error_logs.append(f"TIMEOUT on {url}")
                    row_write[index_bjdt_id].value = ""
                except Exception as e:
                    logger.error(f"ERROR on {url}: {e}")
                    error_logs.append(f"ERROR on {url}: {e}")
                    row_write[index_bjdt_id].value = ""
        
        # Save workbook
        wb_write.save(pengujian_path)
        wb_write.close()
        wb_data.close()
        
        # Save error logs
        if error_logs:
            with open("bjdt_error_log.txt", mode="w", encoding="utf-8") as error_file:
                for log in error_logs:
                    error_file.write(log + "\n")
            logger.warning(f"Saved {len(error_logs)} errors to bjdt_error_log.txt")
        
        logger.info("BJDT ID process completed")
        
        return BjdtResult(
            success=True,
            message=f"Fetched {ids_found} IDs, {ids_missing} missing, {len(error_logs)} errors",
            ids_found=ids_found,
            ids_missing=ids_missing,
            errors=error_logs
        )
        
    except Exception as e:
        logger.error(f"Error during BJDT ID fetch: {str(e)}")
        return BjdtResult(
            success=False,
            message=f"Failed: {str(e)}",
            errors=[str(e)]
        )
